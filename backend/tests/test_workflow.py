import io
import os
import tempfile

os.environ["DATABASE_URL"] = "sqlite:///" + tempfile.mktemp(suffix=".db")
os.environ["APP_SECRET"] = "test-secret-of-more-than-thirty-two-characters"
os.environ["UPLOAD_DIR"] = tempfile.mkdtemp(prefix="ticketing-uploads-")

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from alembic import command
from alembic.config import Config
from sqlalchemy import create_engine, inspect, select, text
from sqlalchemy.dialects import mysql, postgresql
from sqlalchemy.schema import CreateTable

from app.auth import key_digest
import app.db as database
from app.db import Base, BootstrapKey, SessionLocal
from app.main import app
from app.seed_stelinacs import seed as seed_stelinacs


def assert_ok(response):
    assert response.status_code == 200, response.text
    return response.json()


def test_schema_compiles_for_postgres_and_mysql():
    for dialect in (postgresql.dialect(), mysql.dialect()):
        statements = [str(CreateTable(table).compile(dialect=dialect)) for table in Base.metadata.sorted_tables]
        assert len(statements) == 10
        assert all("CREATE TABLE" in statement for statement in statements)


def test_migration_backfills_existing_tickets(monkeypatch, tmp_path):
    legacy_url = "sqlite:///" + str(tmp_path / "legacy.db").replace("\\", "/")
    monkeypatch.setattr(database, "url", legacy_url)
    config = Config("alembic.ini")
    command.upgrade(config, "5972fdbc4062")
    connection = create_engine(legacy_url)
    with connection.begin() as db:
        timestamp = "2026-01-01 00:00:00"
        db.execute(text("INSERT INTO tenants (id, name, slug, created_at, updated_at) VALUES (1, 'One', 'one', :ts, :ts)"), {"ts": timestamp})
        db.execute(text("INSERT INTO users (id, email, password_hash, role, tenant_id, created_at, updated_at) VALUES (1, 'a@b.com', 'hash', 'tenant', 1, :ts, :ts)"), {"ts": timestamp})
        db.execute(text("INSERT INTO steps (id, tenant_id, name, description, created_at, updated_at) VALUES (1, 1, 'Intake', '', :ts, :ts)"), {"ts": timestamp})
        db.execute(text("INSERT INTO tickets (id, tenant_id, step_id, title, status, created_by, created_at, updated_at) VALUES (1, 1, 1, 'Legacy company', 'open', 1, :ts, :ts)"), {"ts": timestamp})
        db.execute(text("INSERT INTO ticket_events (id, tenant_id, ticket_id, actor_id, kind, to_step_id, snapshot, created_at, updated_at) VALUES (1, 1, 1, 1, 'created', 1, '{}', :ts, :ts)"), {"ts": timestamp})
    command.upgrade(config, "head")
    with connection.connect() as db:
        row = db.execute(text("SELECT identifier, title, title_overridden FROM tickets WHERE id=1")).one()
        assert tuple(row) == ("Legacy company", "Legacy company", 0)
        assert db.execute(text("SELECT identifier_label FROM tenants WHERE id=1")).scalar_one() == "Company name"
        assert db.execute(text("SELECT identifier FROM ticket_events WHERE id=1")).scalar_one_or_none() is None
        assert "uploaded_files" in inspect(connection).get_table_names()
    connection.dispose()


def test_tenant_workflow_and_filtered_export():
    command.upgrade(Config("alembic.ini"), "head")
    with SessionLocal.begin() as db:
        db.add(BootstrapKey(key_hash=key_digest("first-key")))
    with TestClient(app) as admin:
        assert admin.post("/api/auth/bootstrap", headers={"Origin": "https://other.example"}, json={"key": "first-key", "email": "admin@example.com", "password": "long-password-123"}).status_code == 403
        assert_ok(admin.post("/api/auth/bootstrap", json={"key": "first-key", "email": "admin@example.com", "password": "long-password-123"}))
        assert admin.post("/api/auth/bootstrap", json={"key": "first-key", "email": "second@example.com", "password": "long-password-123"}).status_code == 409
        one = assert_ok(admin.post("/api/admin/tenants", json={"name": "One", "slug": "one"}))
        two = assert_ok(admin.post("/api/admin/tenants", json={"name": "Two", "slug": "two"}))
        assert len(one["api_key"]) >= 32 and len(two["api_key"]) >= 32
        for tenant, email in ((one, "one@example.com"), (two, "two@example.com")):
            assert_ok(admin.post("/api/admin/users", json={"tenant_id": tenant["id"], "email": email, "password": "tenant-password-123"}))

    with TestClient(app) as a, TestClient(app) as b:
        assert_ok(a.post("/api/auth/login", json={"email": "one@example.com", "password": "tenant-password-123"}))
        assert_ok(b.post("/api/auth/login", json={"email": "two@example.com", "password": "tenant-password-123"}))
        upload = a.post("/api/files", files={"file": ("proof.txt", b"Acme proof", "text/plain")})
        assert upload.status_code == 201, upload.text
        file_info = upload.json()
        assert "api_key=" in file_info["access_url"]
        assert a.get(file_info["access_url"]).content == b"Acme proof"
        assert b.get(f'/api/files/{file_info["file_id"]}/access', params={"api_key": two["api_key"]}).status_code == 404
        assert a.get(f'/api/files/{file_info["file_id"]}/access', params={"api_key": two["api_key"]}).status_code == 404
        assert b.get(f'/api/files/{file_info["file_id"]}/access', params={"api_key": "invalid-key"}).status_code == 422
        updated_workspace = assert_ok(a.put("/api/workspace", json={"identifier_label": "Nama Perusahaan", "workspace_name": "One Workspace"}))
        assert updated_workspace["identifier_label"] == "Nama Perusahaan"
        assert updated_workspace["workspace_name"] == "One Workspace"
        assert assert_ok(b.get("/api/workspace"))["workspace_name"] == "Two"
        icon_bytes = io.BytesIO()
        Image.new("RGB", (12, 12), (20, 120, 110)).save(icon_bytes, format="PNG")
        icon = a.post("/api/workspace/icon", files={"file": ("brand.png", icon_bytes.getvalue(), "image/png")})
        assert icon.status_code == 201, icon.text
        icon_response = a.get("/api/workspace/icon")
        assert icon_response.headers["content-type"] == "image/webp"
        assert Image.open(io.BytesIO(icon_response.content)).size == (256, 256)
        assert b.get("/api/workspace/icon").status_code == 404
        assert a.delete("/api/workspace/icon").status_code == 200
        assert a.get("/api/workspace/icon").status_code == 404
        assert a.post("/api/workspace/icon", files={"file": ("not-an-image.png", b"invalid", "image/png")}).status_code == 422
        shared = assert_ok(a.post("/api/fields", json={"name": "Customer", "key": "customer", "type": "text"}))
        old = assert_ok(a.post("/api/fields", json={"name": "Old note", "key": "old_note", "type": "text"}))
        new = assert_ok(a.post("/api/fields", json={"name": "Approval", "key": "approval", "type": "boolean"}))
        regions = assert_ok(a.post("/api/fields", json={"name": "Regions", "key": "regions", "type": "array"}))
        file_field = assert_ok(a.post("/api/fields", json={"name": "Proof", "key": "proof", "type": "file"}))
        first = assert_ok(a.post("/api/steps", json={"name": "Intake", "fields": [{"field_id": shared["id"], "required": True}, {"field_id": old["id"]}, {"field_id": regions["id"]}, {"field_id": file_field["id"]}]}))
        second = assert_ok(a.post("/api/steps", json={"name": "Review", "fields": [{"field_id": shared["id"], "required": True}, {"field_id": new["id"], "required": True}]}))
        ticket = assert_ok(a.post("/api/tickets", json={"identifier": "Acme", "step_id": first["id"], "values": {str(shared["id"]): "Alice", str(old["id"]): "Private note", str(regions["id"]): ["North", "West"], str(file_field["id"]): file_info}}))
        assert ticket["events"][0]["actor_name"] == "one@example.com"
        assert a.post("/api/tickets", json={"identifier": "Invalid array", "step_id": first["id"], "values": {str(shared["id"]): "Bob", str(regions["id"]): ["North", 12]}}).status_code == 422
        assert ticket["values"][str(file_field["id"])]["access_url"] == file_info["access_url"]
        assert ticket["title"] == "Acme" and not ticket["title_overridden"]
        assert b.get(f'/api/tickets/{ticket["id"]}').status_code == 404
        assert b.put(f'/api/tickets/{ticket["id"]}/identity', json={"identifier": "Other"}).status_code == 404
        renamed = assert_ok(a.put(f'/api/tickets/{ticket["id"]}/identity', json={"identifier": "Acme Ltd"}))
        assert renamed["title"] == "Acme Ltd" and renamed["events"][-1]["kind"] == "identity_updated"
        custom = assert_ok(a.put(f'/api/tickets/{ticket["id"]}/identity', json={"identifier": "Acme Ltd", "title": "Check request"}))
        assert custom["title_overridden"]
        renamed = assert_ok(a.put(f'/api/tickets/{ticket["id"]}/identity', json={"identifier": "Acme Group", "title": "Check request"}))
        assert renamed["title"] == "Check request"
        assert b.post("/api/steps", json={"name": "Foreign", "fields": [{"field_id": shared["id"]}]}).status_code == 404
        assert a.post(f'/api/tickets/{ticket["id"]}/move', json={"step_id": second["id"]}).status_code == 422
        moved = assert_ok(a.post(f'/api/tickets/{ticket["id"]}/move', json={"step_id": second["id"], "values": {str(new["id"]): False}}))
        assert moved["identifier"] == "Acme Group" and moved["title"] == "Check request"
        assert moved["values"] == {str(shared["id"]): "Alice", str(new["id"]): False}
        assert moved["events"][0]["snapshot"][str(old["id"])] == "Private note"
        assert moved["events"][-1]["from_step_id"] == first["id"]
        returned = assert_ok(a.post(f'/api/tickets/{ticket["id"]}/move', json={"step_id": first["id"]}))
        assert returned["identifier"] == "Acme Group" and str(old["id"]) not in returned["values"]
        moved_again = assert_ok(a.post(f'/api/tickets/{ticket["id"]}/move', json={"step_id": second["id"], "values": {str(new["id"]): True}}))
        assert [e["to_step_id"] for e in moved_again["events"] if e["kind"] in ("created", "moved")] == [first["id"], second["id"], first["id"], second["id"]]
        assert_ok(a.post(f'/api/tickets/{ticket["id"]}/close'))
        assert a.put(f'/api/tickets/{ticket["id"]}/identity', json={"identifier": "Closed"}).status_code == 409
        assert a.post(f'/api/tickets/{ticket["id"]}/move', json={"step_id": first["id"]}).status_code == 409
        reopened = assert_ok(a.post(f'/api/tickets/{ticket["id"]}/reopen'))
        assert reopened["status"] == "open" and reopened["closed_at"] is None
        assert reopened["events"][-1]["kind"] == "reopened"
        assert reopened["events"][-1]["from_step_id"] == second["id"]
        edited_after_reopen = assert_ok(a.put(f'/api/tickets/{ticket["id"]}/values', json={"values": moved_again["values"]}))
        assert edited_after_reopen["events"][-1]["kind"] == "edited"
        assert_ok(a.post(f'/api/tickets/{ticket["id"]}/close'))
        reopened_again = assert_ok(a.post(f'/api/tickets/{ticket["id"]}/reopen'))
        assert reopened_again["status"] == "open" and reopened_again["events"][-1]["kind"] == "reopened"
        assert_ok(a.post("/api/tickets", json={"identifier": "Acme Group", "title": "Start in review", "step_id": second["id"], "values": {str(shared["id"]): "Bob", str(new["id"]): True}}))
        intake = assert_ok(a.post("/api/tickets", json={"identifier": "Intake Co", "title": "Still in intake", "step_id": first["id"], "values": {str(shared["id"]): "Chris", str(regions["id"]): ["North", "West"], str(file_field["id"]): file_info}}))
        automatic = assert_ok(a.put(f'/api/tickets/{intake["id"]}/identity', json={"identifier": "Intake Ltd", "title": None}))
        assert automatic["title"] == "Intake Ltd" and not automatic["title_overridden"]
        assert assert_ok(a.get("/api/tickets", params={"step_ids": second["id"], "limit": 1}))["total"] == 2
        assert assert_ok(a.get("/api/tickets", params={"step_ids": first["id"]}))["total"] == 1
        assert assert_ok(a.get("/api/tickets", params={"search": "Acme Group"}))["total"] == 2
        response = a.get("/api/tickets/export", params={"step_ids": second["id"], "search": "Acme Group"})
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet.max_row == 3
        assert sheet["B1"].value == "Nama Perusahaan"
        assert {sheet["B2"].value, sheet["B3"].value} == {"Acme Group"}
        assert {sheet["C2"].value, sheet["C3"].value} == {"Check request", "Start in review"}
        assert "Old note" not in [c.value for c in sheet[1]]
        intake_export = a.get("/api/tickets/export", params={"step_ids": first["id"], "search": "Intake"},
                              headers={"host": "tickets.stelina.example"})
        intake_sheet = load_workbook(io.BytesIO(intake_export.content)).active
        assert "North, West" in [c.value for row in intake_sheet.iter_rows() for c in row]
        exported_file_link = next(c.value for row in intake_sheet.iter_rows() for c in row
                                  if isinstance(c.value, str) and "/api/files/" in c.value)
        assert exported_file_link.startswith("http://tickets.stelina.example/api/files/")
        assert "api_key=" in exported_file_link
        analytics = assert_ok(a.get("/api/analytics/tenant", params={"days": 7}))
        assert analytics["range"]["days"] == 7 and len(analytics["events"]) == 7
        assert analytics["metrics"]["created_tickets"] == 3
        assert {row["step_name"] for row in analytics["tickets_by_step"]} == {"Intake", "Review"}
        assert assert_ok(b.get("/api/analytics/tenant", params={"days": 7}))["metrics"]["created_tickets"] == 0
        assert a.get("/api/admin/analytics", params={"days": 7}).status_code == 403
        assert_ok(a.delete(f'/api/tickets/{intake["id"]}'))
        assert a.post(f'/api/tickets/{intake["id"]}/reopen').status_code == 404
        assert assert_ok(b.get("/api/tickets"))["total"] == 0

    with TestClient(app) as admin:
        assert_ok(admin.post("/api/auth/login", json={"email": "admin@example.com", "password": "long-password-123"}))
        platform = assert_ok(admin.get("/api/admin/analytics", params={"days": 30}))
        assert len(platform["ticket_events"]) == 30 and len(platform["account_growth"]) == 30
        assert set(platform) == {"range", "metrics", "ticket_events", "account_growth", "insights"}
        assert not {"tenant_id", "tenant_name", "name", "slug", "email", "step_name", "identifier", "title", "values"} & set(platform["metrics"])
        assert all(set(point) == {"date", "created", "closed", "reopened", "moved"} for point in platform["ticket_events"])
        assert all(set(point) == {"date", "new_tenants", "new_users"} for point in platform["account_growth"])
        assert "PT Diamond Cold Storage" not in repr(platform)
        tenant_record = next(row for row in assert_ok(admin.get("/api/admin/tenants"))["items"] if row["id"] == one["id"])
        assert "workspace_name" not in tenant_record and "workspace_icon_url" not in tenant_record
        rotated = assert_ok(admin.post(f'/api/admin/tenants/{one["id"]}/api-key'))["api_key"]
    with TestClient(app) as c:
        assert c.get(file_info["access_url"]).status_code == 401
        assert c.get(f'/api/files/{file_info["file_id"]}/access', params={"api_key": rotated}).content == b"Acme proof"


def test_stelinacs_seed_is_idempotent():
    from app.db import MasterField, Step, StepField, Tenant, Ticket, TicketEvent, TicketValue, User
    from app.auth import hash_password
    import json

    with SessionLocal.begin() as db:
        legacy_tenant = Tenant(name="StelinaCS", slug="stelinacs", identifier_label="Nama Perusahaan")
        db.add(legacy_tenant)
        db.flush()
        legacy_user = User(tenant_id=legacy_tenant.id, email="stelinacs1@stelina.co.id",
                           password_hash=hash_password("stelinacs2026"), role="tenant")
        db.add(legacy_user)
        db.flush()
        legacy_fields = {}
        for name, key, kind, options in (("Jenis Kendala", "jenis_kendala", "select", ["Kendala PIB belum tersedia"]),
                                         ("NIB", "nib", "text", []), ("No Aju", "no_aju", "text", []),
                                         ("No PIB", "no_pib", "array", [])):
            field = MasterField(tenant_id=legacy_tenant.id, name=name, key=key, type=kind, options=json.dumps(options))
            db.add(field)
            db.flush()
            legacy_fields[key] = field
        legacy_step = Step(tenant_id=legacy_tenant.id, name="Pelaporan Kendala", description="Legacy single step")
        db.add(legacy_step)
        db.flush()
        for pos, field in enumerate(legacy_fields.values()):
            db.add(StepField(tenant_id=legacy_tenant.id, step_id=legacy_step.id, field_id=field.id,
                             required=field.key == "jenis_kendala", position=pos))
        legacy_ticket = Ticket(tenant_id=legacy_tenant.id, step_id=legacy_step.id,
                               identifier="PT Diamond Cold Storage", title="Kendala PIB belum tersedia",
                               title_overridden=True, status="open", created_by=legacy_user.id)
        db.add(legacy_ticket)
        db.flush()
        for key, value in (("jenis_kendala", "Kendala PIB belum tersedia"), ("nib", "8120008960533"),
                           ("no_aju", "400009B0B0471"), ("no_pib", ["282210", "341296"])):
            db.add(TicketValue(tenant_id=legacy_tenant.id, ticket_id=legacy_ticket.id,
                               field_id=legacy_fields[key].id, value=json.dumps(value)))

    first = seed_stelinacs("stelinacs2026")
    assert first["tickets_created"] == 30
    assert first["tickets_existing"] == 1
    with SessionLocal.begin() as db:
        from app.db import Ticket, now
        ticket = db.scalar(select(Ticket).where(Ticket.identifier == "PT Diamond Cold Storage"))
        assert ticket is not None
        ticket.status = "closed"
        ticket.closed_at = now()
    second = seed_stelinacs("stelinacs2026")
    assert second["tickets_created"] == 0
    assert second["tickets_existing"] == 31
    with SessionLocal() as db:
        from sqlalchemy import func
        tenant = db.scalar(select(Tenant).where(Tenant.slug == "stelinacs"))
        assert tenant is not None and tenant.identifier_label == "Nama Perusahaan"
        assert db.scalar(select(func.count(Ticket.id)).where(Ticket.tenant_id == tenant.id)) == 31
        issue_steps = db.scalars(select(Step).where(Step.tenant_id == tenant.id, Step.deleted_at.is_(None))).all()
        assert len(issue_steps) == 6
        seeded_tickets = db.scalars(select(Ticket).where(Ticket.tenant_id == tenant.id, Ticket.deleted_at.is_(None))).all()
        assert {step.name for step in issue_steps} == {ticket.title for ticket in seeded_tickets}
        assert db.scalar(select(func.count(MasterField.id)).where(MasterField.tenant_id == tenant.id,
                       MasterField.key == "jenis_kendala", MasterField.deleted_at.is_(None))) == 0
        assert db.scalar(select(func.count(Step.id)).where(
            Step.tenant_id == tenant.id, Step.name == "Pelaporan Kendala", Step.deleted_at.is_(None))) == 0
        assert db.scalar(select(func.count(TicketEvent.id)).where(
            TicketEvent.ticket_id == legacy_ticket.id, TicketEvent.kind == "moved")) == 1
        for issue_step in issue_steps:
            assert db.scalar(select(func.count(Ticket.id)).where(
                Ticket.tenant_id == tenant.id, Ticket.step_id == issue_step.id,
                Ticket.title == issue_step.name, Ticket.deleted_at.is_(None))) > 0
        assert db.scalar(select(Ticket.status).where(Ticket.identifier == "PT Diamond Cold Storage")) == "closed"
        user = db.scalar(select(User).where(User.email == "stelinacs1@stelina.co.id"))
        assert user is not None and user.tenant_id == tenant.id
